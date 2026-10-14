"""Quarter export: XLSX with three sheets and CSV, with hours equal to the stored schedule."""

import csv
import io
from collections import defaultdict
from datetime import date, datetime
from typing import Any

import pytest
from openpyxl import load_workbook

from app.exporters.model import text
from tests.api_support import ApiHarness
from tests.test_api_task_management import _objective, _quarter, _reschedule, _task

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMULA_TITLE = '=HYPERLINK("http://example.com")'


def _setup(api: ApiHarness) -> tuple[int, dict[str, Any]]:
    quarter_id = _quarter(api)
    _, kr_high = _objective(api, quarter_id, "high")
    _, kr_low = _objective(api, quarter_id, "low")
    first = _task(api, kr_high, 480)
    _task(api, kr_high, 240, [first])
    risky = _task(api, kr_high, 120)
    api.client.patch(f"/tasks/{risky}", json={"title": FORMULA_TITLE})
    pinned = _task(api, kr_low, 600)
    api.client.put(f"/tasks/{pinned}/pin", json={"start_date": "2026-10-12"})
    return quarter_id, _reschedule(api, quarter_id)["schedule"]


def _export(api: ApiHarness, quarter_id: int, **params: Any) -> Any:
    response = api.client.get(f"/quarters/{quarter_id}/export", params=params)
    assert response.status_code == 200, response.text
    return response


def _as_date(value: Any) -> date | None:
    return value.date() if isinstance(value, datetime) else value


def _rows(sheet: Any) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def test_xlsx_has_three_sheets_whose_hours_match_the_schedule(api: ApiHarness) -> None:
    quarter_id, schedule = _setup(api)
    response = _export(api, quarter_id)
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == 'attachment; filename="okr-q4-2026.xlsx"'

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Công việc", "Lịch theo ngày", "Workload theo tuần"]
    tasks_sheet, days_sheet, weeks_sheet = (workbook[name] for name in workbook.sheetnames)

    header, *task_rows = _rows(tasks_sheet)
    assert header[:5] == ("ID", "Mục tiêu", "Kết quả then chốt", "Công việc", "Ước lượng (giờ)")
    placed = {t["task_id"]: t for t in schedule["tasks"]}
    assert {row[0] for row in task_rows} == set(placed)
    for row in task_rows:
        entry = placed[row[0]]
        assert row[9] == sum(a["minutes"] for a in entry["allocations"]) / 60
        assert (_as_date(row[7]), _as_date(row[8])) == (
            date.fromisoformat(entry["start_date"]),
            date.fromisoformat(entry["end_date"]),
        )
    assert "'" + FORMULA_TITLE in {row[3] for row in task_rows}

    header, *day_rows = _rows(days_sheet)
    assert header[0] == "Ngày"
    assert [(_as_date(r[0]).isoformat(), r[2], r[3]) for r in day_rows if r[0] is not None] == [
        (d["date"], d["capacity_minutes"] / 60, d["scheduled_minutes"] / 60)
        for d in schedule["days"]
    ]
    total = sum(d["scheduled_minutes"] for d in schedule["days"]) / 60
    assert total == sum(t["estimate_minutes"] for t in _all_tasks(api, quarter_id)) / 60

    header, *week_rows = _rows(weeks_sheet)
    assert header[:7] == (
        "Tuần",
        "Bắt đầu",
        "Kết thúc",
        "Ngày làm việc",
        "Có (giờ)",
        "Đã xếp (giờ)",
        "Mức sử dụng (%)",
    )
    assert header[7:] == ("Objective high", "Objective low")
    assert week_rows[0][:6] == ("T40", datetime(2026, 10, 1), datetime(2026, 10, 4), 2, 8, 8)
    assert sum(r[5] for r in week_rows) == total
    per_objective: dict[int, float] = defaultdict(float)
    for entry in schedule["tasks"]:
        per_objective[entry["objective_id"]] += sum(a["minutes"] for a in entry["allocations"])
    assert [sum(r[7] for r in week_rows), sum(r[8] for r in week_rows)] == [
        v / 60 for v in per_objective.values()
    ]
    assert tasks_sheet.freeze_panes == "A2"


def _all_tasks(api: ApiHarness, quarter_id: int) -> list[dict[str, Any]]:
    detail = api.client.get(f"/quarters/{quarter_id}").json()
    return [t for o in detail["objectives"] for kr in o["key_results"] for t in kr["tasks"]]


def test_csv_is_utf8_with_bom_in_the_chosen_language(api: ApiHarness) -> None:
    quarter_id, schedule = _setup(api)
    response = _export(api, quarter_id, format="csv", lang="en", details=True)
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert response.headers["content-disposition"].endswith('filename="okr-q4-2026-tasks.csv"')
    assert response.content.startswith(b"\xef\xbb\xbf")
    assert b"\r\n" in response.content

    header, *rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert header[:4] == ["ID", "Objective", "Key result", "Task"]
    assert header[-2:] == ["Done criteria", "Depends on"]
    assert len(rows) == len(schedule["tasks"])
    by_title = {row[3]: row for row in rows}
    assert "'" + FORMULA_TITLE in by_title
    second = next(row for row in rows if row[-1])
    assert int(second[-1]) in {int(row[0]) for row in rows}

    days = _export(api, quarter_id, format="csv", sheet="days")
    day_header, *day_rows = list(csv.reader(io.StringIO(days.content.decode("utf-8-sig"))))
    assert day_header[0] == "Ngày"
    assert [(r[0], float(r[3])) for r in day_rows] == [
        (d["date"], d["scheduled_minutes"] / 60) for d in schedule["days"]
    ]


def test_export_without_a_schedule_lists_tasks_unscheduled(api: ApiHarness) -> None:
    quarter_id = _quarter(api)
    _, kr_id = _objective(api, quarter_id)
    _task(api, kr_id, 90)
    workbook = load_workbook(io.BytesIO(_export(api, quarter_id).content))
    tasks_sheet, days_sheet, weeks_sheet = (workbook[name] for name in workbook.sheetnames)
    (_, row) = _rows(tasks_sheet)
    assert (row[4], row[9], row[7]) == (1.5, 0, None)
    assert len(_rows(days_sheet)) == 1 and len(_rows(weeks_sheet)) == 1


def test_export_errors(api: ApiHarness) -> None:
    missing = api.client.get("/quarters/999/export")
    assert missing.status_code == 404
    quarter_id = _quarter(api)
    assert api.client.get(f"/quarters/{quarter_id}/export?format=pdf").status_code == 422


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r"])
def test_cells_that_look_like_formulas_are_quoted(prefix: str) -> None:
    assert text(f"{prefix}1+1") == f"'{prefix}1+1"
    assert text("Plain 1+1") == "Plain 1+1"


# --- dondon SQL -----------------------------------------------------------------------------

DONDON = {"email": "me@example.com", "tag_name": "OKR Quý 4"}


def _dondon(api: ApiHarness, quarter_id: int, **body: Any) -> Any:
    return api.client.post(f"/quarters/{quarter_id}/export/dondon-sql", json={**DONDON, **body})


def test_dondon_sql_is_one_guarded_transaction_with_a_row_per_task(api: ApiHarness) -> None:
    quarter_id, schedule = _setup(api)
    response = _dondon(api, quarter_id)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/sql")
    assert "okr-q4-2026-dondon.sql" in response.headers["content-disposition"]
    sql = response.text
    assert sql.index("START TRANSACTION;") < sql.index("COMMIT;")
    assert "WHERE email = 'me@example.com'" in sql
    assert sql.count("INSERT INTO tasks") == len(schedule["tasks"])
    assert sql.count("INSERT INTO task_tags") == len(schedule["tasks"])
    assert "'OKR Quý 4'" in sql
    assert "list_type = 'task'" in sql


def test_dondon_sql_converts_dates_to_utc_and_maps_priority(api: ApiHarness) -> None:
    quarter_id, _ = _setup(api)
    sql = _dondon(api, quarter_id, work_start="09:30", utc_offset_minutes=420).text
    # The pinned low-priority task starts on 2026-10-12 09:30 local = 02:30 UTC.
    assert "'2026-10-12 02:30:00'" in sql
    # Date-only due dates are local midnight, i.e. 17:00 UTC the day before.
    assert " 17:00:00'" in sql
    assert ", 3, 0, 'todo'" in sql  # high priority objective, not done
    assert ", 1, 0, 'todo'" in sql  # low priority objective
    flat = _dondon(api, quarter_id, map_priority=False).text
    assert ", 3, 0, 'todo'" not in flat


def test_dondon_sql_keeps_or_resets_the_status(api: ApiHarness) -> None:
    quarter_id, _ = _setup(api)
    task = api.client.get(f"/quarters/{quarter_id}/tasks").json()[0]["id"]
    assert api.client.patch(f"/tasks/{task}", json={"status": "done"}).status_code == 200
    assert "'done'" in _dondon(api, quarter_id).text
    assert "'done'" not in _dondon(api, quarter_id, keep_status=False).text


def test_dondon_sql_escapes_hostile_text(api: ApiHarness) -> None:
    quarter_id = _quarter(api)
    _, kr = _objective(api, quarter_id)
    task = _task(api, kr, 60)
    nasty = "x'); DROP TABLE users; -- \\ \nline\x00"
    assert api.client.patch(f"/tasks/{task}", json={"title": nasty[:60]}).status_code == 200
    sql = _dondon(api, quarter_id, tag_name="it's \\").text
    assert "DROP TABLE users; --" in sql  # present only as text inside a string literal
    assert "'x''); DROP TABLE users;" in sql
    assert "\x00" not in sql and "\0" not in sql
    assert "'it''s \\\\'" in sql


@pytest.mark.parametrize(
    "body",
    [
        {"email": "not-an-email"},
        {"email": "a'b@example.com"},
        {"email": "a@b.com; DROP TABLE users"},
        {"tag_name": "   "},
        {"work_start": "25:00"},
        {"utc_offset_minutes": 900},
        {"extra": 1},
    ],
)
def test_dondon_sql_rejects_bad_input(api: ApiHarness, body: dict[str, Any]) -> None:
    quarter_id, _ = _setup(api)
    assert _dondon(api, quarter_id, **body).status_code == 422


def test_dondon_sql_for_a_missing_quarter_is_404(api: ApiHarness) -> None:
    assert _dondon(api, 999).status_code == 404


def test_dondon_sql_writes_ids_as_32_hex_characters(api: ApiHarness) -> None:
    """The dondon app binds ids without dashes; rows written with UUID() stay invisible to it."""
    quarter_id, schedule = _setup(api)
    sql = _dondon(api, quarter_id).text
    new_id = "REPLACE(UUID(), '-', '')"
    assert sql.count(new_id) == 1 + 2 * len(schedule["tasks"])  # tag, then task and link per task
    assert sql.count("UUID()") == sql.count(new_id)
