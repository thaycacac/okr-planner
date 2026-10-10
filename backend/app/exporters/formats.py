"""CSV and XLSX adapters of the exporter port."""

import csv
import io
from datetime import date
from typing import Protocol

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.exporters.model import SHEETS, Cell, ExportData, SheetName, tables

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
_MAX_COLUMN_WIDTH = 60


class Exporter(Protocol):
    extension: str
    media_type: str

    def render(self, data: ExportData) -> bytes: ...


class CsvExporter:
    """One table per file, UTF-8 with BOM so Excel reads Vietnamese correctly; ISO dates."""

    extension = "csv"
    media_type = CSV_MEDIA_TYPE

    def __init__(self, sheet: SheetName = "tasks") -> None:
        self.sheet = sheet

    def render(self, data: ExportData) -> bytes:
        (table,) = tables(data, [self.sheet])
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\r\n")
        writer.writerow(table.headers)
        writer.writerows([_csv_cell(c) for c in row] for row in table.rows)
        return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def _csv_cell(cell: Cell) -> str | int | float:
    if cell is None:
        return ""
    if isinstance(cell, date):
        return cell.isoformat()
    return cell


class XlsxExporter:
    """One workbook with the three sheets: tasks, daily schedule, weekly workload."""

    extension = "xlsx"
    media_type = XLSX_MEDIA_TYPE

    def render(self, data: ExportData) -> bytes:
        workbook = Workbook()
        default = workbook.active
        if default is not None:
            workbook.remove(default)
        bold = Font(bold=True)
        for table in tables(data, SHEETS):
            sheet = workbook.create_sheet(table.title)
            sheet.append(list(table.headers))
            for cell in sheet[1]:
                cell.font = bold
            widths = [len(h) for h in table.headers]
            for row in table.rows:
                sheet.append(list(row))
                for index, value in enumerate(row):
                    shown = "" if value is None else str(value)
                    widths[index] = max(widths[index], len(shown))
            for row_cells in sheet.iter_rows(min_row=2):
                for cell in row_cells:
                    if isinstance(cell.value, date):
                        cell.number_format = "yyyy-mm-dd"
            for index, width in enumerate(widths, start=1):
                letter = get_column_letter(index)
                sheet.column_dimensions[letter].width = min(width + 2, _MAX_COLUMN_WIDTH)
            sheet.freeze_panes = "A2"
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
